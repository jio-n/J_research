"""Create a separate narrative Word report and a report-oriented HTML reader."""
import html,json,pathlib,re
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.section import WD_SECTION,WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=pathlib.Path(__file__).resolve().parents[1]
W=R/'.build/technical_report';W.mkdir(exist_ok=True)
OUT=R/'outputs';STEM='JIMTOF2026_技術動向レポート_20261019'
D=json.loads((R/'data/research/report_v0/report.json').read_text('utf-8'))
CASES={c['card_id']:c for t in D['topics'] for c in t['cases']}
EV={e['evidence_id']:e for e in D['evidence']}
GROUPS=[
 ('S01','図面と原価情報を使う見積支援',['FC062']),
 ('S02','加工履歴に基づく工程予測',['FC018']),
 ('S03','CAMの加工戦略と経路の作成支援',['FC020','AC131']),
 ('S04','PLCとHMIの設計支援',['FC011','FC012']),
 ('S05','設備アラームの診断と対処案',['FC013']),
 ('S06','電気設計の対話支援',['FC017']),
 ('S07','NC検証と切削物理による最適化',['FC022']),
 ('S08','主軸負荷を使う加工中の適応制御',['RPT002']),
 ('S09','寸法計測と工具補正の連携',['FC033']),
 ('S10','工具刃先の3D計測と補正加工',['AC085']),
 ('S11','外観検査と追加学習',['FC036']),
 ('S12','工具摩耗の判定と寿命の推定',['FC039','FC040']),
 ('S13','工具画像を使う摩耗診断',['FC038']),
 ('S14','稼働データを使う設備診断',['AC225']),
 ('S15','2024年にも存在した加工異常検知',['FC064']),
 ('S16','主軸に絡む切粉の認識と除去',['FC049']),
 ('S17','3D認識による面取り軌道の生成',['AC146']),
 ('S18','小部品の認識と磁気ピック',['AC051']),
 ('S19','バリ取り工具の摩耗への自動追従',['AC097']),
 ('S20','爪と工具の自動交換',['AC214']),
 ('S21','移設後のロボット位置復旧',['FC048']),
 ('S22','学習モデルの作成と実装支援',['FC014']),
 ('S23','設備情報とAIエージェントの接続',['FC024']),
 ('S24','工作機械の仮想検証',['FC007']),
 ('S25','工作機械ツインと工程内品質保証',['AC098']),
 ('S26','物理シミュレーション用データの生成',['AC273']),
 ('S27','工具データと在庫及び購買の連携',['AC295']),
 ('S28','搬送車の遠隔操作による復帰',['FC031']),
 ('S29','双腕移動ヒューマノイドの遠隔教示',['RPT005']),
 ('S30','加工前の歯車計測と開発中の最適化',['RPT004']),
 ('S31','ワークとパレットの供給自動化',['FC002']),
 ('S32','熱変位制御とびびり抑制',['FC009','RPT007']),
]
SPECIAL={'S04':['ERF-01','EDP-03'],'S05':['ERF-02','EAT-N04'],'S11':['ERF-05','EDP-14'],'S12':['EOV-W08','EDP-12'],'S16':['ERF-13'],'S22':['EDP-04'],'S25':['EDP-16','EDP-17'],'S26':['ERF-06','ERF-09']}
REFS=[]
for sid,title,ids in GROUPS:
 es=[EV[e]for k in ids for e in CASES[k]['evidence_ids']];chosen=[]
 for k in ids:
  for year in [2026,2024]:
   e=next((EV[i]for i in CASES[k]['evidence_ids'] if EV[i]['year']==year and 'jimtof.org/'in EV[i]['url'] and 'ESdetails'in EV[i]['url']),None)
   if e and e['url'] not in [x['url']for x in chosen]:chosen.append(e)
 extra=[EV[i]for i in SPECIAL.get(sid,[]) if i in EV]
 if not extra:extra=[e for e in es if 'jimtof.org/'not in e['url']][:1]
 for e in extra:
  if e['url'] not in [x['url']for x in chosen]:chosen.append(e)
 REFS.append(dict(id=sid,title=title,card_ids=ids,companies=list(dict.fromkeys(CASES[k]['company']for k in ids)),evidence=chosen))
REFBY={x['id']:x for x in REFS};SID={k:sid for sid,title,ids in GROUPS for k in ids}

# The report's prose explains the relationships; it is not slide copy.
CHAPTERS=[
 dict(id='preparation',title='生産準備と制御設計',lead='図面を探す、加工順序を考える、制御コードを書くといった準備作業に、AIによる案の生成が入り始めている。狙いは、毎回ゼロから作業を組み立てる負担を減らし、経験の浅い担当者でも熟練者に確認できるたたき台を早く作ることにある。',sections=[
 ('見積と加工準備で扱う知識が異なる','【公式記載】匠フォースは図面認識、類似案件参照、原価計算から見積書案を作る。ProPlanAIは形状特徴と過去ジョブを使って加工工程を予測し、CAMAssistは工具・材料・機械等の条件から加工戦略や経路の案を生成する。[S01–S03]'),
 ('案を作った後の確認まで見る','【分析】これらは同じ「AIによる自動化」でも、使う知識と出力が異なる。見積では原価の根拠、工程検討では保持や工具の制約、CAMでは加工性とNCの整合を確かめる必要がある。作成時間だけでなく、人が直す量と初品までの手戻りを含めて比較したい。'),
 ('制御設計への組込みが進む','【公式記載】TwinCAT CoAgent for EngineeringはPLCコードやI/O構成、HMIの設計を支援する。Eplan Copilotは設計文脈に応じて回答・推奨・手順を作る。前者は専門家による確認と承認、後者も人による意思決定を前提とする。[S04,S06] 【未確認】作成されたコードが実機に書き込まれ、運転を開始するまでの権限や検証の範囲は、個別に確認が必要である。')],cards=['FC062','FC018','FC020','FC011','FC017'],insight='【年度比較】PLCコード支援は2023年、CAMAssistのMastercam対応は2024年にも公表されている。2026年の特徴は、AIの誕生そのものより、支援する業務や設計環境への組込みが具体化している点にある。[S03,S04]'),
 dict(id='machining',title='加工条件の最適化と品質の自動補正',lead='加工時間を短くしたい一方で、過負荷や熱変位、寸法のばらつきは避けたい。この要求に対し、加工前の予測、加工中の状態監視、計測後の補正という異なる時点で技術が使われている。',sections=[
 ('いつ判断して何を変えるかが方式を分ける','【公式記載】Vericut Forceは切削物理と力・動力の制約を使い、加工前にNCの送りを最適化する。ACMは加工中の主軸負荷を監視し、その場で送りを調整する。Renishaw CentralとIPCは計測結果を工具オフセット等へ戻す。これらを一括して学習AIと扱うことはできない。[S07–S09]'),
 ('計測と加工を近づける提案','【公式記載】サイバーRCのGH5は機内3D画像で刃先を測り、幅・角度を狙い値へ近づけるAI補正ホーニングを説明する。目標値の設定は人が担い、AIの内部方式や測定不能時の復帰は未公表である。[S10] 【分析】測ってから加工へ戻す構成は、機外測定や再段取りの負担を減らす方向として注目できる。'),
 ('既存制御との関係を確かめる','【年度比較】Force、負荷適応制御、Centralは2024年にも説明されている。2026年に見られる計測と加工の一体化は、既存の補正を置き換える新原理と決めつけず、適用対象と構成の拡張として比較する。[S07–S10] 【未確認】材料・工具を替えた際の設定工数、補正上限、異常値を反映させない条件が残る。')],cards=['FC022','RPT002','FC033','AC085'],insight='【分析】事前予測、加工中の適応制御、計測後の補正は相互に補える。ただし、別製品がそのまま一つにつながり、品質を保証するとは確認していない。熱変位制御やびびり抑制も、AIの有無を分けて評価する必要がある。[S32]'),
 dict(id='inspection',title='外観検査と工具及び設備の状態判断',lead='外観、刃先、設備の振動など、人が目や経験で見極めていた対象に画像や稼働データが使われる。判定を自動化するだけでなく、その結果を継続使用、交換、保全の判断へどうつなぐかが比較の焦点になる。',sections=[
 ('検査を続けながら判定を調整する','【公式記載】HACARUS Check ZEROは良品との差で異常を検出し、人の良否注釈を使って追加調整する。学習画像や注釈の準備は人の仕事であり、更新モデルの採用や出荷判定の権限は別に確認する必要がある。追加学習は2025年に発表済みで、2026年初の機能とは扱わない。[S11]'),
 ('摩耗の判定と交換時期の提案は別の仕事','【公式記載】モリマシナリーはツールコラージュの摩耗・欠け判定に加え、アナライザーで摩耗データから寿命と交換時期を提案する。ハイデンハインはVT122とVTC 6.0によるAI摩耗診断を記載する。[S12,S13] 【分析】画像上の摩耗量が分かることと、特定の加工で使い続けてよいことは同じではない。加工条件と品質の履歴を合わせて比較する必要がある。'),
 ('設備診断にも既存技術がある','【公式記載】ムラテックCCSは電流・振動等を新品状態と比較して劣化や異常を診断する。三菱電機も2024年にAI・機械学習による加工異常、工具摩耗、作業ミスの検知を説明していた。[S14,S15] 【未確認】対象工具・設備を変えた際の精度、誤警報への対処、停止・交換の最終決定者は十分に比較できていない。')],cards=['FC036','FC039','FC040','FC038','AC225'],insight='【分析】判定から次の行動へ支援を広げる提案が見られる。ただし、寿命予測には以前の研究発表もあり、判定から工具交換までの自動実行を全社が提供するという意味ではない。[S12]'),
 dict(id='physical',title='段取りと付帯作業に広がる自動化',lead='ワークを運ぶ部分を自動化しても、切粉の絡み、工具の摩耗、治具の交換、再教示に人手が残ると長時間運転は続きにくい。認識と動作を組み合わせる技術と、交換機構や追従機構で解決する技術を、同じ作業の負担から比較することが重要である。',sections=[
 ('何を認識してどの作業を行うか','【公式記載】TAKISAWAは主軸の切粉の絡まりを判断して除去する参考出品を記載する。東京エレクトロンデバイスは3D認識から不定形ワークの面取り・皿もみ軌道を生成する。カネテックはカメラ、協働ロボット、磁気ピンセットによる小部品ピックにAIを明記している。[S16–S18] 【未確認】各社でAIが担う具体的な判断、失敗検出、再試行、完了確認までの範囲は異なり、未公開の点も多い。'),
 ('機構による解決も比較対象になる','【公式記載】XEBEC CPSオートホルダーはブラシ摩耗に応じて突出し量を調整する。松本機械工業は爪・工具等の交換を、中村留精密工業は移設後のマーカーによる位置復旧を説明する。[S19–S21] 【分析】これらは、AIの導入によらず人の調整や段取りを減らせる対照技術である。'),
 ('機体の形と自律性を分ける','【公式記載】東京貿易テクノシステムのRBY Seriesは遠隔操作・教示の参考出展であり、自律作業の範囲は未確認である。[S29] 【年度比較】搬入出や爪交換、小部品ピックは2024にも存在する。2026は切粉処理や品種変更に伴う準備など、残る作業へ対象を広げる提案が具体化している。[S16,S18,S20,S31]')],cards=['FC049','AC146','AC051','AC097','AC214'],insight='【開発仮説】汎用人型機体の開発を前提にするより、切粉除去やバリ取りなど対象を絞り、認識・工具機構・動作・完了判断を組み合わせる領域に検討価値がある。実用性は成功率と再教示時間、人の介入回数で確かめる。'),
 dict(id='operation',title='工場運用支援とAIの実装基盤',lead='設備、工具、進捗、保全記録が別々にあると、予定変更や停止のたびに人が状況を集め直すことになる。データを見えるようにする技術に加え、集めた情報から原因や対処を提案するエージェントが具体化している。',sections=[
 ('状況把握と提案と実行を分ける','【公式記載】TwinCAT CoAgent for Operationsは機械データと文書から原因仮説・対処案を作る。専門家が制御権を持つ方針だが、書込み権限や毎回の承認、提供時期には確認が残る。WinToolは在庫ルールに基づく購買要求を管理者承認後にERPへ送る。後者をAIによる自律発注とする根拠はない。[S05,S27]'),
 ('例外への対処には人の運用も残る','【公式記載】X-Area Remoteの実証では、停止車両を人が遠隔操作して回避する。AI支援型という呼称があっても、確認できた例外操作の主体は人である。[S28] 【分析】提案が出ることと、現場でその提案を実行できることを分け、停止時に誰が判断し、どこから復帰するかを比較したい。'),
 ('接続と学習と仮想検証を用途に結び付ける','【公式記載】たけびしはOPC UA情報モデルとエージェントの連携、ベッコフはAutoMLとONNXによるモデル出力、ファナックは加工前の仮想検証を説明する。NdotLightのTRINIXはCADから形状・物理属性を生成し、検証や自己修正をうたう。[S22–S26] 【未確認】生成属性と実機の一致や、自己修正の収束は確認できていない。シーメンスの海外ヒューマノイド実証とJIMTOFの展示を、同じ統合システムとして扱わない。[S25,S26]')],cards=['FC013','AC295','FC024','FC014','AC273'],insight='【分析】基盤の価値は、生産準備、加工、ロボット等の具体的な仕事をどれだけ助けるかで見る。データ接続やツインの導入自体は、設備の自律運転を示さない。')]

MATRIX=[
 ['役割','受注と見積','加工準備','制御設計','段取りと付帯作業','加工と調整','検査と工具','工場運用'],
 ['集めて見せる','匠フォース\n図面認識 [S01]','Vericut\nNCの仮想確認 [S07]','代表例なし','RoboSync\nマーカー認識 [S21]','GH5\n刃先3D計測 [S10]','モリマシナリー\n刃先撮像 [S12]','ムラテックCCS\n電流と振動 [S14]'],
 ['状態を見極める','代表例なし','Vericut\n干渉等の検出 [S07]','代表例なし','TAKISAWA\n切粉絡み判定 [S16]','ACM\n負荷の監視 [S08]','VTC 6.0\nAI摩耗診断 [S13]','プライマリドクタ\n劣化診断 [S14]'],
 ['案や手順を作る','匠フォース\n見積書案 [S01]','ProPlanAIとCAMAssist\n工程と経路案 [S02,S03]','CoAgent\nPLCとHMI案 [S04]','東京ED\n面取り軌道 [S17]','Force\nNC送り最適化 [S07]','モリマシナリー\n交換時期提案 [S12]','CoAgent Operations\n対処案 [S05]'],
 ['作業や調整を行う','送信権限等は未確認','実機実行とは別','起動権限等は未確認','MMK 爪交換\nTAK 切粉除去は参考 [S16,S20]','ACM 送り調整\nIPC 補正 [S08,S09]','WinTool\n承認後ERP送信 [S27]','X-Area Remote\n人が遠隔操作 [S28]'],
]

def fmt(run,size=11,bold=False,color='000000'):
 run.font.name='Arial';run.font.size=Pt(size);run.bold=bold;run.font.color.rgb=RGBColor.from_string(color)
 rf=run._element.get_or_add_rPr().get_or_add_rFonts();rf.set(qn('w:eastAsia'),'Meiryo');rf.set(qn('w:ascii'),'Arial');rf.set(qn('w:hAnsi'),'Arial')
 return run
doc=Document()
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
 st=doc.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0);st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Meiryo')
 for border in list(st._element.get_or_add_pPr().findall(qn('w:pBdr'))):border.getparent().remove(border)
doc.styles['Normal'].font.size=Pt(11);doc.styles['Normal'].paragraph_format.line_spacing=Pt(16.5);doc.styles['Normal'].paragraph_format.space_after=Pt(6)
snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');doc.styles['Normal']._element.get_or_add_pPr().append(snap)
for name,size in [('Title',26),('Subtitle',12),('Heading 1',18),('Heading 2',12.5),('Heading 3',11)]:
 st=doc.styles[name];st.font.size=Pt(size);st.paragraph_format.space_after=Pt(7);st.paragraph_format.space_before=Pt(10 if name=='Heading 2' else 0)
 st.paragraph_format.keep_with_next=True
def setup(sec,land=False):
 sec.orientation=WD_ORIENT.LANDSCAPE if land else WD_ORIENT.PORTRAIT
 sec.page_width=Inches(11 if land else 8.5);sec.page_height=Inches(8.5 if land else 11)
 sec.top_margin=Inches(.62);sec.bottom_margin=Inches(.58);sec.left_margin=sec.right_margin=Inches(.65);sec.header_distance=Inches(.22);sec.footer_distance=Inches(.25)
 grid=sec._sectPr.find(qn('w:docGrid'))
 if grid is not None:sec._sectPr.remove(grid)
setup(doc.sections[0]);header=doc.sections[0].header.paragraphs[0];fmt(header.add_run('JIMTOF2026 技術動向調査'),9)
footer=doc.sections[0].footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
fmt(footer.add_run('2026年10月19日　事前調査　'),9)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
landscape=False
def page(title,land=False):
 global landscape
 if land!=landscape:setup(doc.add_section(WD_SECTION.NEW_PAGE),land);landscape=land
 heading=doc.add_heading(title,1)
 if land==landscape:heading.paragraph_format.page_break_before=True
def p(text,style=None,size=None):
 a=doc.add_paragraph(style=style);fmt(a.add_run(text),size or 11);return a
def h(text):doc.add_heading(text,2)
def link(par,url,label,size=9.3):
 rid=par.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
 a=OxmlElement('w:hyperlink');a.set(qn('r:id'),rid);rr=OxmlElement('w:r');pr=OxmlElement('w:rPr')
 for tag,val in [('color','225B92'),('u','single'),('sz',str(round(size*2)))]:
  e=OxmlElement('w:'+tag);e.set(qn('w:val'),val);pr.append(e)
 rf=OxmlElement('w:rFonts');rf.set(qn('w:eastAsia'),'Meiryo');rf.set(qn('w:ascii'),'Arial');pr.append(rf);rr.append(pr)
 t=OxmlElement('w:t');t.text=label;rr.append(t);a.append(rr);par._p.append(a)
def table(headers,rows,widths,size=9.5):
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for col,w in zip(t.columns,widths):col.width=Inches(w)
 for values in [headers]+rows:
  cells=t.rows[0].cells if values is headers else t.add_row().cells
  for i,(cell,txt,w) in enumerate(zip(cells,values,widths)):
   cell.width=Inches(w);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;pr=cell._tc.get_or_add_tcPr()
   margins=OxmlElement('w:tcMar')
   for side in ['top','bottom','left','right']:
    v=OxmlElement('w:'+side);v.set(qn('w:w'),'90');v.set(qn('w:type'),'dxa');margins.append(v)
   pr.append(margins);sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'244B63' if values is headers else ('F0F4F7' if len(t.rows)%2==0 else 'FFFFFF'));pr.append(sh)
   pp=cell.paragraphs[0];pp.paragraph_format.space_after=Pt(2);pp.paragraph_format.space_before=Pt(1);pp.paragraph_format.line_spacing=Pt(size*1.35)
   fmt(pp.add_run(str(txt)),size,values is headers,'FFFFFF' if values is headers else '000000')
   if i==0 and len(headers)>4:pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
  cant=OxmlElement('w:cantSplit');t.rows[-1]._tr.get_or_add_trPr().append(cant)
 pr=t._tbl.tblPr;borders=OxmlElement('w:tblBorders')
 for side in ['top','left','bottom','right','insideH','insideV']:
  e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
 pr.append(borders);repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
 doc.add_paragraph().paragraph_format.space_after=Pt(2)
 return t

doc.add_paragraph('JIMTOF2026\n技術動向調査レポート','Title')
doc.add_paragraph('AIとDXと自動化が担う仕事と研究開発への示唆','Subtitle')
p('2026年10月19日　事前調査',size=10)
h('要旨')
p('【分析】公開されている展示予定と企業の技術資料からは、生産準備の案づくり、加工条件の調整、検査・診断、段取り・付帯作業、工場運用の判断支援へ、AIと自動化の対象が広がる様子が読み取れる。共通する狙いは、現場に残る人手と技能への依存を減らすことである。')
p('2024年以前にも、CAM支援、負荷適応制御、工具摩耗判定、爪交換、デジタルツインは存在した。今回の比較では、単純な新旧の区分より、対象作業、システムの構成、人の担当、提供段階の違いが重要となる。2026年の展示予定を、実演や導入効果が確認された事実とは扱わない。')
table(['調査対象','2024年','2026年'],[['公式一覧の掲載行','1,281','1,146'],['説明本文を保存した掲載行','988','820']],[4.0,1.6,1.6],10)
p('詳細URLのない619掲載行は内容未取得。掲載行は企業数ではなく、公式発表総数との突合は未完了。関連候補索引は344カード・309企業名キー、詳細比較は57事例であり、同じ深さの調査ではない。',size=9.5)
h('本報告の主な示唆')
for v in ['生産準備では、図面・工程・設備の知識を扱う案づくりと、人の確認・修正の組合せが焦点となる。','加工では、物理モデル、実測による適応制御、計測フィードバックを組み合わせる余地がある。','物理作業では、切粉除去やバリ取りなど、変動する付帯作業に対象を絞った開発が候補となる。','工場運用では、原因や代替案の提示に加え、既存システムの情報と実行権限をそろえる必要がある。']:p(v)
p('本文は全体マップ、分野別の比較、年度変化、開発候補、現地確認、出典の順に構成する。「公式記載」はメーカー説明、「分析」はAIによる横断整理、「開発仮説」は未検証の提案、「未確認」は根拠が不足する範囲を表す。',size=9.5)

page('1 技術全体像と役割の分担',True)
p('【分析】同じAIという名称でも、情報を集める、状態を見極める、案を作る、作業を実行する役割は異なる。下表は代表事例の配置で、企業間の優劣や自律性の順位を示さない。',size=10.5)
table(MATRIX[0],MATRIX[1:],[1.0]+[1.24]*7,9)
p('代表例なしは今回の選定範囲の意味であり、対応する企業や需要が存在しないことを示さない。別々の製品・機能を一つの統合システムとして配置したものではない。切粉除去は参考出品、遠隔操作は人が実行する。',size=9.5)
h('工程を横断する基盤')
p('設備情報の接続、学習モデルの作成、デジタルツイン、物理シミュレーション用データは複数の仕事を支える。OPC UA、AutoML、仮想検証等を、どの仕事の準備や判断に使うかと合わせて評価する。[S22–S26]',size=10.5)

for n,ch in enumerate(CHAPTERS,2):
 page(str(n)+' '+ch['title'])
 p('【分析】'+ch['lead'])
 for title,body in ch['sections'][:2]:h(title);p(body)
 # Repeated records belong in a compact comparative table, not narrative boxes.
 rows=[]
 for k in ch['cards']:
  c=CASES[k];rows.append([c['company']+'\n'+c['item']+' ['+SID[k]+']',c['official_mechanism'],c['human_boundary']])
 table(['企業と技術','何を助けるか 公式説明の要約','人の仕事と未確認事項'],rows,[2.1,2.65,2.45],9.3)
 title,body=ch['sections'][2];h(title);p(body)
 p(ch['insight'],size=10)

page('7 2024年から2026年に見える変化')
p('【分析】個別企業の比較を束ねると、既存技術の継続に加え、支援する作業の具体化や、判断結果を次の作業へ結び付ける提案が見られる。ただし、これは公開記載の比較であり、市場全体の普及や性能向上を示すものではない。')
year_rows=[
 ['生産準備','PLC作成支援やCAMAssistのMastercam対応は既存。','見積の一連の処理、工程予測、I/OやHMI等の支援を具体化。','対象業務と組込み先の広がり。[S01–S04]'],
 ['加工条件と補正','Force、ACM、Centralが2024にも掲載。','GH5など計測と加工の一体化を具体記載。','基本原理の継続と構成の拡張。[S07–S10]'],
 ['検査と工具診断','AI摩耗判定や加工異常検知は2024にも存在。','VTC 6.0のAI診断、寿命・交換提案等を具体化。','判定後の判断支援にも注目。ただし以前の研究あり。[S11–S15]'],
 ['段取りと物理作業','爪交換、搬入出、小部品ピックが既存。','切粉除去の参考出品、面取り軌道生成、位置復旧を記載。','対象作業・構成の広がり。AIの担当は製品別。[S16–S21]'],
 ['工場運用','工具・CAM・在庫連携や遠隔支援は既存。','機械データと文書を参照する運用エージェントを発表。','提案から実行への権限・提供段階は未確認。[S05,S27,S28]'],
 ['共通基盤','AutoML、機器接続、工作機械ツインは既存。','AI連携、物理属性を付けた学習用データ等の用途を具体化。','実機との一致は別途確認が必要。[S22–S26]'],
]
table(['技術トピック','2024年以前の基準点','2025年から2026年の記載','差分の読み方'],year_rows,[1.1,1.9,2.2,2.0],9.5)
h('新登場と記載の具体化を区別する')
p('企業の説明が詳しくなったことは、新しい発明の証拠にはならない。追加学習は2025年、PLC支援は2023年等の発表を含めて見る。2024年の同じ製品資料がない事例は比較不能とし、記載減も機能の撤退とは判断しない。')
p('【未確認】各方式の対象範囲、実演の成否、顧客での稼働状況、定量的な工数・品質の改善は、展示予定の説明だけではそろわない。FD100の計測情報による条件最適化のように、開発中の記載も区別して扱う。[S30]')

page('8 研究開発の候補と狙う仕事')
p('以下は公開情報を横断したAIの分析と開発仮説である。複数社が提案する仕事を起点にした候補であり、市場の空白や自社の投資適合性を実証したものではない。')
research_copy=[
 ('特定の部品群で、図面と類似工程を検索し、設備・工具・保持の制約を踏まえた工程案を作る。既存CAMとNC検証を使い、人の修正を次回の準備へ反映する。','生成AIやCAMそのものの再開発より、加工知識と制約を扱う部分を検討する。既存製品だけの運用と、総準備時間・修正量・初品の手戻りを比較する。[S01–S04]'),
 ('一つの加工法に絞り、切削予測、加工中の信号、工具画像、寸法計測から、加工継続・条件変更・工具交換の候補を決める。','適応制御、IPC、摩耗判定は既存である。個別機能と連携方式を同じ品質条件で比べ、工具使用量、加工時間、条件変更時の再調整を評価する。[S07–S15]'),
 ('切粉除去やバリ取り等で、画像・3D認識と工具機構を組み合わせ、作業位置や軌道を決める。作業後の状態を確認し、再動作か人の介入を選ぶ。','既存の認識・力追従・自動交換を比較対象とする。汎用人型機体を前提にせず、異なるワークでの完了率、教示時間、介入回数、タクトを見る。[S16–S21]'),
 ('工具不足や設備停止など一つの例外業務で、設備・工具・進捗の状態と過去の対処を使い、代替案と担当作業を提示する。承認後の結果も追跡する。','既存MES、工具管理、遠隔支援と比較する。利用可能なAPIと権限が得られる範囲に限定し、状況確認時間、実行できた提案、停止時間、入力負担を測る。[S05,S23,S27,S28]'),
]
for i,(t,(approach,comparison))in enumerate(zip(D['research'],research_copy),1):
 h(str(i)+' '+t['name'])
 p('【開発仮説】'+approach)
 p('【既存技術との比較】'+comparison,size=10)
h('共通基盤の扱い')
p('機器接続や仮想環境を独立した大規模システムとして先に整えるのではなく、対象とする部品群・加工法・付帯作業・例外業務を選び、その仕事に必要な範囲で取り込む。既存製品だけで負担を十分に減らせる場合は、連携や適用を優先する。')

page('9 現地で深める確認と報告の更新')
p('事前資料で分かる機能の有無から、実際に使える条件へ確認を進める。訪問先は技術方式が異なる企業を組み合わせ、単一製品の説明だけで研究テーマを決めない。下表は重点比較の論点であり、会場の巡回順を指定するものではない。')
field_rows=[
 ['生産準備と設計','匠技研、Hexagon、JBM、ベッコフ','同じ部品・要求で、入力条件、生成案、人の修正、最終採用までを見せてもらう。'],
 ['加工条件と補正','CGTech、西部商工、レニショー、サイバーRC','材料・工具が変わった時の再設定、補正上限、誤計測時の扱いを確認する。'],
 ['検査と診断','HACARUS、モリマシナリー、ハイデンハイン、ムラテックCCS','対象外工具、不良の見逃しと過検出、学習更新、交換・保全の決定者を聞く。'],
 ['物理作業','TAKISAWA、東京ED、カネテック、MMK、XEBEC','失敗や絡まり残りへの対応、完了判定、教示時間と人の介入を確かめる。'],
 ['運用支援と基盤','ベッコフ、たけびし、WinTool、NdotLight','取得・書込可能な情報、実行許可、仮想と実機の差、提供開始とサポート範囲を聞く。'],
]
table(['確認する分野','比較する訪問候補','見たい説明や証拠'],field_rows,[1.2,2.3,3.7],10)
h('技術と事業の両方を聞く')
p('技術面では、同じ対象条件での作業時間、初品品質、工具使用量、品種変更時の設定、人の介入を確認する。事業面では、有償導入・PoC・引き合いを区別し、対象顧客、導入工数、費用、受注や稼働実績を聞く。市場規模が提示された場合は、対象範囲と算定根拠も記録する。')
h('現地確認後に更新する判断')
p('まず、対象の仕事と提供段階が明確で、既存方式との違いを説明できるかを確認する。次に、顧客が負担と感じる作業が残り、評価に使えるデータが得られるかを確かめる。その結果から、既存製品の適用・連携で進める範囲と、研究開発を要する範囲を分ける。')
p('比較表と研究仮説は現地情報で更新する。宣伝上の呼称、実演、顧客での稼働、将来構想を区別し、現地で確認できなかった点は未確認のまま残す。')

for part,start in enumerate(range(0,len(REFS),11),1):
 page('出典資料 '+str(part))
 if part==1:p('本文のS番号に対応する公式資料。JIMTOFの該当箇所は原則として「出展紹介」「主な出展物」。各社の説明・主張を根拠とし、独立した性能評価とは区別する。資料確認は2026年10月16日から18日、横断整理は19日。',size=9.5)
 for ref in REFS[start:start+11]:
  a=doc.add_paragraph();a.paragraph_format.keep_with_next=True;a.paragraph_format.space_before=Pt(5);a.paragraph_format.space_after=Pt(2);a.paragraph_format.line_spacing=Pt(12.5);fmt(a.add_run(ref['id']+'　'+'／'.join(ref['companies'])+'　'+ref['title']),9.3,True)
  a=doc.add_paragraph();a.paragraph_format.keep_with_next=True;a.paragraph_format.space_after=Pt(2);a.paragraph_format.line_spacing=Pt(12)
  for j,e in enumerate(ref['evidence']):
   if j:fmt(a.add_run('　／　'),9)
   label=str(e['year'])+'出展説明' if 'jimtof.org/'in e['url']and'ESdetails'in e['url'] else '公式技術資料'+str(j+1)
   link(a,e['url'],label,9.3)
  loc=[e['locator'] for e in ref['evidence'] if 'jimtof.org/'not in e['url']]
  text='技術資料の該当箇所　'+' ／ '.join(x[:78] for x in loc) if loc else '該当箇所　出展紹介と主な出展物'
  a=p(text,size=9);a.paragraph_format.line_spacing=Pt(12);a.paragraph_format.space_after=Pt(3)
doc.core_properties.title='JIMTOF2026 技術動向調査レポート';doc.core_properties.subject='AI DX 自動化 自律化の技術動向と研究開発への示唆';doc.core_properties.author='技術調査';doc.core_properties.keywords='JIMTOF2026,事前調査,技術動向'
doc.save(OUT/(STEM+'.docx'))
CONTENT=dict(title='JIMTOF2026 技術動向レポート',as_of='2026-10-19',scope=D['catalog']['scope'],chapters=CHAPTERS,research=D['research'],references=REFS,matrix=MATRIX,year_rows=year_rows,field_rows=field_rows,cases=[{**{k:v for k,v in CASES[k].items() if k!='full_record'},'source_group':SID[k]} for k in SID])
(W/'report_content.json').write_text(json.dumps(CONTENT,ensure_ascii=False,indent=2),'utf-8')
print(json.dumps(dict(docx=str(OUT/(STEM+'.docx')),references=len(REFS),selected_case_rows=len(CONTENT['cases'])),ensure_ascii=False))

H=lambda t:html.escape(str(t))
style='''*{box-sizing:border-box}html{scroll-behavior:smooth}body{margin:0;background:#f4f6f8;color:#182d43;font:16px/1.85 Meiryo,sans-serif}header{background:#183b56;color:white;padding:48px max(30px,calc((100vw - 1450px)/2)) 35px}header h1{font-size:38px;line-height:1.5;margin:8px 0 16px}header p{max-width:1100px}a{color:#225b92}header a{color:white}nav{position:sticky;top:0;background:white;border-bottom:1px solid #cbd6e0;display:flex;gap:23px;padding:13px 30px;overflow:auto;z-index:3}nav a{white-space:nowrap;text-decoration:none;font-weight:bold}main{max-width:1510px;margin:auto;padding:30px}section{scroll-margin-top:80px;margin-bottom:45px}h2{font-size:28px;line-height:1.5}h3{font-size:21px}.intro{font-size:20px;max-width:1200px}.grid{display:grid;grid-template-columns:repeat(2,minmax(0,1fr));gap:20px}.brief{background:white;padding:24px;border:1px solid #d2dce5}.brief h3{margin-top:0;color:#225b92}.meta{font-size:13px;color:#617587}.badge{font-size:12px;padding:3px 7px;background:#e7eef5;color:#425e77}.analysis{background:#fff0d6;color:#7e551c}.scroll{overflow:auto}table{border-collapse:collapse;width:100%;background:white;font-size:14px}th,td{padding:13px;border:1px solid #d5dfe7;vertical-align:top}thead{background:#244b63;color:white}tbody tr:nth-child(2n){background:#f0f4f7}.matrix{min-width:1400px;table-layout:fixed;font-size:12px}.matrix th:first-child{width:125px}.matrix td{white-space:pre-line}.toolbar{display:flex;gap:12px;flex-wrap:wrap;margin:20px 0}button,input,select{font:inherit;padding:8px 12px;border:1px solid #afbfcd;border-radius:3px;background:white}button{cursor:pointer}button.active{background:#225b92;color:white}.case-list{display:grid;grid-template-columns:repeat(3,minmax(0,1fr));gap:18px}.case{background:white;border:1px solid #c8d6e2;padding:22px}.case h3{margin:0 0 6px}.case p{font-size:14px;margin:11px 0}.case .target{font-size:18px;font-weight:bold}.links{display:flex;gap:10px;flex-wrap:wrap;font-size:13px}.year-columns{display:grid;grid-template-columns:1fr 1fr;gap:16px}.year-columns div{background:white;padding:22px;border-top:4px solid #8da9bf}.year-columns div+div{border-color:#225b92}details{margin-top:12px}summary{cursor:pointer;color:#225b92}.source{background:white;padding:15px 22px;border-bottom:1px solid #d2dce5}.source p{font-size:13px;margin:4px 0}.source a{overflow-wrap:anywhere}[hidden]{display:none!important}.tabs{display:flex;gap:9px;flex-wrap:wrap}footer{border-top:1px solid #cbd6e0;padding-top:20px;color:#617587;font-size:13px}@media(max-width:1000px){.case-list{grid-template-columns:repeat(2,minmax(0,1fr))}}@media(max-width:700px){main{padding:20px 16px}header{padding:30px 18px}.grid,.case-list,.year-columns{grid-template-columns:1fr}header h1{font-size:28px}nav{gap:18px;padding:10px 16px}}'''
parts=['<!doctype html><html lang="ja"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>'+H(CONTENT['title'])+'</title><style>'+style+'</style><header><p>2026年10月19日　公式資料に基づく事前調査</p><h1>JIMTOF2026 技術動向レポート</h1><p>現場のどの仕事を、どの技術が助けるのか。生産準備、加工、検査、物理作業、工場運用を横断し、2024年との違いと開発の候補を読む。</p><p><a href="'+STEM+'.pptx">スライド版</a>　<a href="'+STEM+'.docx">Word版</a>　<a href="JIMTOF2026_技術整理と全体マップ_v1.html">詳細な企業索引とマップ</a></p></header><nav>'+''.join(f'<a href="#{i}">{v}</a>'for i,v in [('summary','主要発見'),('overview','仕事の全体像'),('topics','分野別動向'),('comparison','企業事例'),('change','年度の変化'),('proposal','開発と現地確認'),('sources','公式出典')])+'</nav><main>']
parts.append('<section id="summary"><h2>現場に残る準備と判断と付帯作業に注目する</h2><p class="intro">既存技術の継続に加え、支援する仕事を広げる提案が見られます。AIという名称を入口にせず、対象の仕事、判断・実行の範囲、人の関わり、提供段階を比べます。</p><div class="grid">')
summaries=[('準備を助ける','図面・原価・過去工程を使い、見積や加工工程、制御設計の案を作る。確認と修正を含めて省力効果を見る。','preparation'),('加工を調整する','切削物理の予測、加工中の負荷適応、寸法計測からの補正。判断する時点と使う入力が異なる。','machining'),('状態を見極める','外観や工具摩耗、設備の劣化を判定し、継続使用・交換・保全の判断へつなぐ。','inspection'),('変わる物理作業に対応する','切粉や不定形ワークを認識し、除去や面取りへ使う。交換・追従の機構による解決とも比較する。','physical')]
for title,body,anchor in summaries:parts.append(f'<article class="brief"><span class="badge analysis">横断分析</span><h3>{title}</h3><p>{body}</p><a href="#topic-{anchor}">分野の根拠を見る →</a></article>')
parts.append('</div><p class="meta">2026年は展示予定。説明本文は2024年988行・2026年820行を保存。詳細URLなし619掲載行は内容未取得。関連索引344カードと詳細比較57事例の調査深度は異なります。市場シェア・普及率を示す集計ではありません。</p></section>')
parts.append('<section id="overview"><h2>仕事と担当する役割の全体像</h2><p>代表事例の配置〈分析〉。空欄・代表例なしは、技術や需要の不在を示しません。別々の展示を一つの統合システムとは扱いません。</p><div class="scroll"><table class="matrix"><thead><tr>'+''.join('<th>'+H(x)+'</th>'for x in MATRIX[0])+'</tr></thead><tbody>'+''.join('<tr>'+''.join('<td>'+H(x)+'</td>'for x in row)+'</tr>'for row in MATRIX[1:])+'</tbody></table></div><p class="meta">横断基盤：機器接続、学習モデル作成、デジタルツイン、物理シミュレーション用データ。用途ごとの必要性で評価します。</p></section>')
parts.append('<section id="topics"><h2>分野ごとに見たい技術の違い</h2>')
html_briefs=[('準備する情報と生成する案が異なる','見積は図面と原価、工程予測は過去ジョブ、CAMは工具・材料・機械の条件を扱います。案を作る速さに加え、担当者の確認と修正がどれだけ残るかを見ます。'),('加工前と加工中と計測後を分ける','NC送りの事前最適化、実測負荷に応じた送り調整、計測値によるオフセット補正は異なる方式です。複数をつなぐ場合も、補正の限界や異常値の扱いが必要です。'),('判定と交換判断を分ける','画像で摩耗が分かっても、加工を続けてよいかは別の判断です。良品学習の更新や寿命推定を、対象工具と加工条件に結び付けて比較します。'),('搬送以外に残る仕事を見る','切粉除去や面取り、ブラシ摩耗への追従、爪交換が焦点です。AI認識、3D幾何認識、機構による解決を同じ作業から比較し、完了確認と失敗時の引継ぎを確かめます。'),('提案するAIと実行する運用をつなぐ','設備データと文書からの対処案、承認後の購買、遠隔での復帰を区別します。接続や学習・仮想環境は、こうした具体業務の実装に必要な範囲で評価します。')]
for ch,(title,body)in zip(CHAPTERS,html_briefs):
 parts.append(f'<article id="topic-{ch["id"]}" class="brief"><span class="badge analysis">動向の解釈</span><h3>{H(ch["title"])}</h3><p><b>{title}</b><br>{body}</p><div class="links">'+''.join(f'<a href="#sources-{SID[k]}">{H(CASES[k]["company"])}の原典</a>'for k in ch['cards'])+f'</div><button type="button" data-focus-topic="{ch["id"]}">この分野の企業事例を比較</button></article>')
parts.append('</section><section id="comparison"><h2>具体的な展示と人の役割を比較する</h2><div class="toolbar"><select id="report-topic"><option value="ALL">全分野</option>'+''.join(f'<option value="{c["id"]}">{H(c["title"])}</option>'for c in CHAPTERS)+'</select><input id="report-search" placeholder="企業・技術・ブースを検索" aria-label="企業事例を検索"></div><p class="meta" id="report-count"></p><div class="case-list">')
for c in CONTENT['cases']:
 topic_sets={'preparation':['B01','B02','B04'],'machining':['B03','B05','B06'],'inspection':['B07','B08','B09'],'physical':['B10','B11','B12','B17'],'operation':['B13','B14','B15','B16']}
 ts=[key for key,ids in topic_sets.items() if c['case_id'].split('-')[0]in ids];ref=REFBY[c['source_group']]
 parts.append(f'<article class="case" data-topics="{H(" ".join(ts))}" data-search="{H((c["company"]+c["item"]+c["target"]+c["booth"]).lower())}"><h3>{H(c["company"])}</h3><span class="meta">{H(c["item"])}　{H(c["booth"])}</span><p class="target">{H(c["target"])}</p><p><span class="badge">公式説明の要約</span><br>{H(c["official_mechanism"])}</p><details><summary>人の仕事と2024年との差分</summary><p><b>人・未確認</b><br>{H(c["human_boundary"])}</p><p><b>年度差</b><br>{H(c["year_change"])}</p><p>{H(c["ai_status"])}</p></details><div class="links">'+''.join(f'<a href="{H(e["url"])}" target="_blank" rel="noopener">{e["year"]} '+('出展説明' if 'jimtof.org/'in e['url']and'ESdetails'in e['url']else'公式技術資料')+' ↗</a>'for e in ref['evidence'])+'</div></article>')
parts.append('</div></section><section id="change"><h2>2024年からの変化を技術トピックで読む</h2><div class="tabs">'+''.join(f'<button type="button" data-year-choice="{i}"'+(' class="active"'if i==0 else'')+'>'+H(row[0])+'</button>'for i,row in enumerate(year_rows))+'</div>')
for i,row in enumerate(year_rows):parts.append(f'<article data-year="{i}"'+(' hidden'if i else'')+f'><div class="year-columns"><div><h3>2024年以前</h3><p>{H(row[1])}</p></div><div><h3>2025年から2026年</h3><p>{H(row[2])}</p></div></div><p><span class="badge analysis">比較分析</span>{H(row[3])}</p></article>')
parts.append('<p class="meta">2024に記載がないことは新発明の証拠ではありません。個別資料不足は比較不能です。参考出品・製品発表・開発中の機能を実用化済みと扱いません。</p></section><section id="proposal"><h2>開発候補と現地で確かめる条件</h2><p>以下は未検証の開発仮説です。市場の空白・優位性・自社の投資適合性を実証したものではありません。</p><div class="grid">')
for t in D['research']:parts.append(f'<article class="brief"><span class="badge analysis">開発仮説</span><h3>{H(t["name"])}</h3><p>{H(t["position"])}</p><details><summary>技術の狙いと比較する効果</summary><p>{H(t["approach"])}</p><p><b>評価</b> {H(t["metrics"])}</p><p>{H(t["decision"])}</p></details></article>')
parts.append('</div><h3>現地では同じ対象条件で比較する</h3><div class="scroll"><table><thead><tr><th>分野</th><th>訪問候補</th><th>確認したい証拠</th></tr></thead><tbody>'+''.join('<tr>'+''.join('<td>'+H(x)+'</td>'for x in row)+'</tr>'for row in field_rows)+'</tbody></table></div><p>採用・受注・PoC・引き合いを区別し、導入工数、対象顧客、費用、市場規模の算定根拠も確認します。結果に応じ、既存製品を適用・連携する範囲と研究開発を要する範囲を分けます。</p></section><section id="sources"><h2>主な公式資料</h2><p class="meta">本文のS番号に対応。JIMTOFの出展紹介・主な出展物、および企業公式の製品・技術資料。メーカー主張と独立した性能検証を区別します。</p>')
for ref in REFS:
 parts.append(f'<article class="source" id="sources-{ref["id"]}"><h3>{ref["id"]}　{H(ref["title"])}</h3>'+''.join(f'<p><a href="{H(e["url"])}" target="_blank" rel="noopener">{H(e["title"])} ↗</a><br>{H(e["evidence_id"])}　確認 {H(e["checked_date"])}<br>該当箇所 {H(e["locator"])}</p>'for e in ref['evidence'])+'</article>')
parts.append('</section><footer>公式記載はメーカーの説明、動向の整理は分析、研究開発候補は仮説です。2026年は展示予定。既存の企業索引HTMLを別途参照できます。</footer></main><script>\nfunction filterReport(){const q=document.getElementById("report-search").value.toLocaleLowerCase(),t=document.getElementById("report-topic").value;let n=0;document.querySelectorAll(".case").forEach(c=>{c.hidden=!(c.dataset.search.includes(q)&&(t==="ALL"||c.dataset.topics.split(" ").includes(t)));if(!c.hidden)n++});document.getElementById("report-count").textContent=n+"事例を表示"}document.getElementById("report-search").oninput=filterReport;document.getElementById("report-topic").onchange=filterReport;document.querySelectorAll("[data-focus-topic]").forEach(b=>b.onclick=()=>{document.getElementById("report-topic").value=b.dataset.focusTopic;document.getElementById("report-search").value="";filterReport();location.hash="comparison"});document.querySelectorAll("[data-year-choice]").forEach(b=>b.onclick=()=>{document.querySelectorAll("[data-year]").forEach(e=>e.hidden=e.dataset.year!==b.dataset.yearChoice);document.querySelectorAll("[data-year-choice]").forEach(e=>e.classList.toggle("active",e===b))});filterReport();\n</script></html>')
(OUT/(STEM+'.html')).write_text(''.join(parts),'utf-8')
print('New report HTML written without editing the existing reader')

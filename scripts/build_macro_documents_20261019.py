import base64
import html
import json
from collections import defaultdict
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
DATA=ROOT/'data/research/macro_20261019'
BUILD=ROOT/'.build/macro_20261019'
OUT=ROOT/'outputs/AI_AX_DX_20261019'
BUILD.mkdir(parents=True,exist_ok=True);OUT.mkdir(parents=True,exist_ok=True)
d=json.loads((DATA/'report_content.json').read_text(encoding='utf-8'))
census=d['census'];sources={s['id']:s for s in d['sources']}
tags=['AI明示','生成AIとエージェント','フィジカルAI明示','自動化と省人化','設計製造のデジタル化','工場運用とデータ連携','状態監視と補正','自律の明示']
labels=['AIの語','生成AI・エージェント','フィジカルAI','自動化・省人化','設計・製造デジタル','工場運用・データ','状態監視・補正','自律の語']

def esc(t):return html.escape(str(t),quote=True)
def refs(ids):
    return ' '.join(f'<a href="{esc(sources[i]["url"])}" target="_blank" rel="noopener">[{esc(i)}]</a>' for i in dict.fromkeys(ids) if i in sources)
def web_table(headers,rows):
    return '<div class="tablewrap"><table><thead><tr>'+''.join('<th>'+esc(x)+'</th>' for x in headers)+'</tr></thead><tbody>'+''.join('<tr>'+''.join('<td>'+str(x)+'</td>' for x in row)+'</tr>' for row in rows)+'</tbody></table></div>'

coverage_rows=[]
for y in ['2024','2026']:
    x=census[y];coverage_rows.append([y,f"{x['listing_rows']:,}",f"{x['name_keys']:,}",f"{x['text_name_keys']:,}",f"{x['text_coverage_pct']}%",str(x['no_text_name_keys'])])
theme_rows=[]
for t in tags:
    a,b=census['2024'],census['2026'];p=census['matched_panel']['counts'][t]
    theme_rows.append([esc(t),f"{a['theme_counts'][t]} / {a['theme_rates'][t]}%",f"{b['theme_counts'][t]} / {b['theme_rates'][t]}%",f"{p['2024']} → {p['2026']}"])
trend_html=''.join(f'<article id="{t["id"]}" class="trend"><span class="kicker">{t["id"]}</span><h3>{esc(t["title"])}</h3><p class="lead">{esc(t["finding"])}</p><div class="compare"><div><h4>2024年以前</h4><p>{esc(t["past"])}</p></div><div><h4>2026年の展示予定と確認資料</h4><p>{esc(t["now"])}</p></div></div><p><strong>仕組みと比較軸</strong>　{esc(t["mechanism"])} {esc(t["implication"])}</p><p class="caveat">{esc(t["limit"])}</p><p><strong>会場での質問</strong>　{esc(t["question"])}</p><p class="refs">{refs(t["refs"])}</p></article>' for t in d['trends'])
major_html=''.join(f'<article class="company" data-group="{esc(m["group"])}" data-search="{esc(m["company"]+m["exhibits"]+m["delta"])}"><div class="companyhead"><h3>{esc(m["name"])}</h3><span>{esc(" / ".join(m["booths"]))}</span></div><p class="meta">{esc(m["group"])}　訪問区分 {m["priority"]}</p><p><strong>見るもの</strong>　{esc(m["exhibits"])}</p><p><strong>2024年との比較</strong>　{esc(m["delta"])}</p><p><strong>確認点</strong>　{esc(m["question"])}</p><p class="refs">{refs(m["refs"])}</p></article>' for m in d['majors'])
visit_html=''.join(f'<article class="visit" data-search="{esc(v["name"]+v["why"]+" ".join(v["booths"]))}"><div class="companyhead"><h3>{v["id"]}　{esc(v["name"])}</h3><span>{esc(" / ".join(v["booths"]))}</span></div><p>{esc(v["why"])}</p><p><strong>確認すること</strong>　{esc(v["ask"])}</p><p class="meta">目安 {v["minutes"]}分（移動・待ち時間別）　{esc(v["status"])}</p><p>{refs(v["refs"])}</p></article>' for v in d['visits'])
route_html=''.join('<h3>'+esc(r['day'])+'</h3>'+web_table(['時間枠','エリア','候補と目的'],[[esc(x) for x in b] for b in r['blocks']]) for r in d['route'])
directory=[dict(name=r['company'],booths=r['booths'],tags=r['tags'],available=bool(r['text']),urls=r['urls']) for r in json.loads((DATA/'corpus.json').read_text(encoding='utf-8')) if r['year']==2026]
source_html=''.join(f'<li id="source-{esc(s["id"])}"><a href="{esc(s["url"])}" target="_blank" rel="noopener">[{esc(s["id"])}] {esc(s["title"])}</a><small>公開 {esc(s["published"])} ／ 確認 {esc(s["checked"])} ／ {esc(s["locator"])}</small></li>' for s in d['sources'])
chartb64=base64.b64encode((BUILD/'panel.png').read_bytes()).decode()
htmltext='''<!doctype html><html lang="ja"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>JIMTOF2026 AIと自動化の調査ガイド</title><style>
:root{--ink:#172d3d;--teal:#126b84;--muted:#536675;--line:#d6e0e6;--pale:#f1f5f7}*{box-sizing:border-box}body{margin:0;font:16px/1.9 "Noto Sans JP","Yu Gothic",Meiryo,sans-serif;color:var(--ink);background:#fff}a{color:#075e80;text-underline-offset:3px}header{background:#15374a;color:white;padding:64px max(6vw,24px) 48px}header h1{font-size:clamp(30px,4vw,52px);line-height:1.4;margin:.3em 0}header p{max-width:930px;color:#d3e3ec}.kicker{font-size:13px;font-weight:700;letter-spacing:.08em}.wrap{max-width:1240px;margin:auto;padding:0 32px}nav{position:sticky;top:0;z-index:5;background:#fff;border-bottom:1px solid var(--line);padding:12px 24px;display:flex;flex-wrap:wrap;gap:12px 24px;font-size:14px}nav a{text-decoration:none}section{padding:48px 0;border-bottom:1px solid var(--line);scroll-margin-top:100px}h2{font-size:28px;margin:0 0 24px}h3{font-size:21px;line-height:1.6;margin:6px 0 16px}h4{margin:0 0 8px}.lead{font-size:19px;font-weight:600;line-height:1.8}.summary{padding-left:1.4em}.summary li{margin:0 0 18px}.metrics{display:grid;grid-template-columns:repeat(3,1fr);gap:24px;margin:28px 0}.metric{border-top:3px solid var(--teal);padding-top:16px}.metric strong{display:block;font-size:36px}.meta,small{font-size:13px;color:var(--muted)}.caveat{color:#76521b;background:#faf6ed;padding:16px 20px}.tablewrap{overflow-x:auto;margin:22px 0}table{border-collapse:collapse;width:100%;min-width:580px;font-size:14px}th,td{padding:13px 15px;border:1px solid var(--line);vertical-align:top}th{text-align:left;background:#e9f0f4}tr:nth-child(even){background:#f8fafb}.trend{padding:32px 0;scroll-margin-top:100px;border-top:1px solid var(--line)}.compare{display:grid;grid-template-columns:1fr 1fr;gap:28px}.compare>div{padding:18px 20px;background:var(--pale)}.refs{font-size:13px;overflow-wrap:anywhere}.company,.visit{padding:24px 0;border-top:1px solid var(--line)}.company p,.visit p{margin:8px 0}.companyhead{display:flex;justify-content:space-between;gap:18px;align-items:baseline}.companyhead span{color:var(--teal);font-weight:700}.filters{display:flex;gap:12px;flex-wrap:wrap;align-items:center;margin:24px 0}input,select,button{font:inherit;border:1px solid #9eb3c0;padding:9px 12px;border-radius:4px;background:white;color:var(--ink)}input{min-width:250px;flex:1}button{cursor:pointer}button:disabled{opacity:.4;cursor:default}input:focus,select:focus,button:focus{outline:3px solid #89c5d7;outline-offset:2px}img.chart{width:100%;height:auto;max-width:970px}.source-list{padding-left:1.5em}.source-list li{margin-bottom:17px;overflow-wrap:anywhere}.source-list small{display:block}footer{padding:32px;background:var(--pale)}[hidden]{display:none!important}@media(max-width:700px){.wrap{padding:0 20px}.compare,.metrics{grid-template-columns:1fr}.companyhead{display:block}nav{position:static}.companyhead h3{margin-bottom:2px}header{padding-top:38px}section{padding:32px 0}}@media print{nav,.filters,button{display:none}section{break-before:auto}header{background:white;color:black;padding:20px}header p{color:black}.company,.visit,.trend{break-inside:avoid}details{display:block}body{font-size:10pt}.wrap{padding:0}}
</style></head><body><header><div class="kicker">JIMTOF2026　事前調査　2026年10月19日</div><h1>AIと自動化の動向<br>展示を比較し 訪問先を決める</h1><p>全出展者一覧の集計と2024年比較を土台に、AX／DX、生成AI、フィジカルAI、長時間無人運転を読み解く。会期は10月26〜31日、東京ビッグサイト。2026年の情報は展示予定で、現地実演と性能は未検証。</p></header><nav aria-label="レポート目次"><a href="#summary">結論</a><a href="#macro">全体と年度比較</a><a href="#trends">6つの動向</a><a href="#majors">大手33社</a><a href="#visits">優先訪問19件</a><a href="#route">回り方</a><a href="#directory">全出展者索引</a><a href="#sources">出典</a></nav><main class="wrap">
'''
htmltext+='<section id="summary"><h2>先に押さえる結論</h2><ol class="summary">'+''.join('<li>'+esc(x)+'</li>' for x in d['findings'])+'</ol><div class="metrics"><div class="metric"><strong>2,427行</strong>両年の全掲載行を処理</div><div class="metric"><strong>538</strong>両年に説明がある同名比較</div><div class="metric"><strong>22 → 40</strong>同名比較でのAI言及</div></div><p>本稿でのAXは、AIによって業務の担当・判断・手順が変わる動きを指す整理上の呼称。DX、自動化、AI利用、自律化は別の軸で比較する。</p></section>'
htmltext+='<section id="macro"><h2>全体集計と2024年からの変化</h2><p>記事や企業を先に数十件選ぶ方法ではなく、保存済み公式一覧の全行から出発した。NFKC・空白除去・英字大小統一で名称キーを作り、同名の複数ブースの説明を統合。法人同一性の完全な名寄せではない。語彙分類は重複する。</p>'+web_table(['年度','掲載行','名称キー','説明あり','説明充足率','説明なし'],coverage_rows)+'<p>母集団集計は2026年10月16日取得のCSVに固定し、最新の企業発表は個別分析へ追加した。異なる取得時点を混ぜて増減率を作らない。</p><img class="chart" src="data:image/png;base64,'+chartb64+'" alt="同名538名称キーでのテーマ言及数。AIは22から40、自動化は172から193。">'+web_table(['テーマの語彙','2024 件数／比率','2026 件数／比率','同名538の比較'],theme_rows)+'<p class="caveat">比率の分母は各年の説明あり933／800名称キー。市場シェア・導入率・技術成熟度ではない。説明なし309／331名称キーは集計で陰性にしない。掲載総数の差は会場構成・掲載定義の影響もあり、産業の縮小を示さない。</p><h3>集計と文脈の照合</h3><p>両年のAI語ヒット86件を文脈確認した。2024のAI/AOはアナログ入出力、2026のAIサーバー向け材料・AI産業向けスピンドル・翻訳などは、製造AIの機能実装の証拠にしない。図の数値は再現可能な語彙集計として保持し、技術分析で除外理由を示す。</p><p>同名比較ではAIの新規言及25件、継続15件、記載が消えたもの7件。追加記載は新機能の初登場を、記載消失は機能の廃止を意味しない。工場運用データは59から57、設計製造デジタルは43から39で、DX全体が一様に増えたとは言えない。</p><p>構造的背景として、2026年版ものづくり白書は、現場データと実装・運用を含むAIロボティクスの政策を整理している。主催者企画でもロボット部品加工ときさげ技能が題材になる。これらは展示の方向性を補強するが、商用導入実績の代わりにはしない。'+refs(['S03','S02','S04'])+'</p></section>'
htmltext+='<section id="trends"><h2>仕事と実現方式から見る6つの動向</h2><p>以下は複数社の説明を束ねた分析。各社が同じ速度・段階で進化しているという意味ではない。</p>'+trend_html+'</section>'
htmltext+='<section id="majors"><h2>大手と主要企業33社の比較</h2><p>Aは今回のAI・自動化の比較を組み立てる優先訪問先、Bは関心工程に応じて追加する比較先。企業の優劣・市場シェアの順位ではない。大手・著名企業を明示的に収録し、AI記載の有無だけで選別しない。</p><div class="filters"><label for="major-q">企業・技術</label><input id="major-q" placeholder="例：工具、計測、DMG"><label for="major-group">分野</label><select id="major-group"><option value="">全分野</option>'+''.join('<option>'+esc(g)+'</option>' for g in dict.fromkeys(m['group'] for m in d['majors']))+'</select></div><div id="major-list">'+major_html+'</div></section>'
htmltext+='<section id="visits"><h2>優先訪問と比較価値</h2><p>全体の調査対象を19件に限定したものではない。主要トレンドごとに、異なる方式・提供段階・残る人手を比較するための訪問候補を選んだ。全19件を1日で回る想定にはしていない。</p>'+visit_html+'</section>'
htmltext+='<section id="route"><h2>2日間の回り方と当日の確認</h2><p>時間枠は配分案で予約や実演時刻ではない。各枠で候補を選び、移動・待ち時間を確保する。東4〜6ホールは改修で休館。全日10〜17時、最終日は16時まで。'+refs(['S01'])+'</p>'+route_html+'<p class="caveat">3社以上を挙げた枠では2社を目安に選択する。DMG森精機の枠は同社を優先し、ジェイテクトは入替候補。南3の候補は2社程度に絞る。最終日は16時に終了するよう調整。</p><h3>1日で見る場合</h3><p>'+esc(d['one_day'])+'</p><h3>主催者企画と時間指定の候補</h3><p>東ガレリアのロボティクス企画は、AIを使う機械と、AIロボットを製造するための加工需要の違いを見る機会。Kaleido開発者との対談は10月26日14〜15時、10月30日11〜12時・14〜15時が掲載されている。南4ではきさげロボットと匠の協業。ミツトヨの品質管理DXセミナーは10月30日11時20分〜12時05分。参加条件・最新時刻は原典で確認。'+refs(['S02','S13'])+'</p>'+web_table(['対象','見る指標','質問'],[[esc(v) for v in x] for x in d['questions']])+'</section>'
htmltext+='<section id="directory"><h2>2026年の全出展者索引</h2><p>全1,131名称キー。テーマは保存CSVの語彙ヒットで、精査済みの機能分類ではない。「語彙ヒットなし」「説明なし」も表示する。公式詳細がない行は公式一覧へのリンクで補う。</p><div class="filters"><label for="dir-q">企業・ブース</label><input id="dir-q" placeholder="企業名または小間番号"><label for="dir-tag">テーマ</label><select id="dir-tag"><option value="">全件</option>'+''.join('<option>'+esc(t)+'</option>' for t in tags)+'<option value="missing">説明なし</option></select></div><p id="dir-count" aria-live="polite"></p><div class="tablewrap"><table><thead><tr><th>企業</th><th>ブース</th><th>語彙候補</th><th>原典</th></tr></thead><tbody id="dir-body"></tbody></table></div><div class="filters"><button id="prev">前へ</button><span id="dir-page"></span><button id="next">次へ</button></div><noscript>索引の検索にはJavaScriptが必要です。本文の分析・企業比較・出典はJavaScriptなしでも読めます。</noscript></section>'
htmltext+='<section id="sources"><h2>出典と調査の範囲</h2><p>本文の[S]は追加照合した公式資料・企業発表、[C]は保存CSVの両年公式詳細、その他のIDは既存の技術別根拠台帳。メーカーの性能主張と調査者の分析を区別し、第三者による性能実証とは扱わない。出典にある国外実証・他展示会の内容は背景資料とし、JIMTOF出展を自動的に認定しない。</p><p>既存整理は334カード・309企業名称を含むが、個別タスクの重点69カードと追加候補は精査の深さが異なる。この件数を全出展者の採用率には使わない。全件語彙集計、AI語86件の文脈確認、大手33社比較、6つの横断分析を今回の成果として分けて示した。</p><details><summary>出典一覧を開く（'+str(len(d['sources']))+'件）</summary><ol class="source-list">'+source_html+'</ol></details><p class="caveat">残る確認課題は、説明なし企業、2024の別名照合、会場での提供版、性能・導入実績・価格、失敗時の復帰。情報がないことを「進化なし」「展示なし」とは判定しない。</p></section></main><footer><div class="wrap">調査基準日 2026年10月19日 ／ 公式カタログ母集団 10月16日取得 ／ ローカルで閲覧可能。原典リンクを開く際はネット接続が必要。</div></footer>'
htmltext+='<script id="directory-data" type="application/json">'+json.dumps(directory,ensure_ascii=False).replace('</',r'<\/')+'</script>'+r'''
<script>
const eq=s=>s.normalize('NFKC').toLocaleLowerCase().replace(/\s/g,'');
const mq=document.getElementById('major-q'),mg=document.getElementById('major-group');function filterMajors(){const q=eq(mq.value);document.querySelectorAll('.company').forEach(c=>c.hidden=!(eq(c.dataset.search).includes(q)&&(!mg.value||c.dataset.group===mg.value)))}mq.addEventListener('input',filterMajors);mg.addEventListener('change',filterMajors);
const ds=JSON.parse(document.getElementById('directory-data').textContent),dq=document.getElementById('dir-q'),dt=document.getElementById('dir-tag'),db=document.getElementById('dir-body');let page=0;const pageSize=40;let found=[];function addCell(tr,value){const td=document.createElement('td');td.textContent=value;tr.append(td);return td}function draw(){const q=eq(dq.value);found=ds.filter(r=>eq(r.name+r.booths.join(' ')).includes(q)&&(!dt.value||(dt.value==='missing'?!r.available:r.tags.includes(dt.value))));const pages=Math.max(1,Math.ceil(found.length/pageSize));page=Math.min(page,pages-1);db.replaceChildren();for(const r of found.slice(page*pageSize,(page+1)*pageSize)){const tr=document.createElement('tr');addCell(tr,r.name);addCell(tr,r.booths.join(' / '));addCell(tr,r.available?(r.tags.join('、')||'語彙ヒットなし'):'説明なし');const td=addCell(tr,'');const urls=r.urls.length?r.urls:['https://www.jimtof.org/jp/exhi_search_pronoun'];urls.forEach((u,i)=>{const a=document.createElement('a');a.href=u;a.target='_blank';a.rel='noopener';a.textContent=r.urls.length?'詳細'+(urls.length>1?i+1:''):'公式一覧';td.append(a,document.createTextNode(' '))});db.append(tr)}document.getElementById('dir-count').textContent=found.length+'件一致';document.getElementById('dir-page').textContent=(page+1)+' / '+pages;document.getElementById('prev').disabled=page===0;document.getElementById('next').disabled=page>=pages-1;}for(const input of[dq,dt])input.addEventListener(input===dq?'input':'change',()=>{page=0;draw()});document.getElementById('prev').onclick=()=>{page--;draw()};document.getElementById('next').onclick=()=>{page++;draw()};draw();
</script></body></html>'''
(OUT/'JIMTOF2026_AI_AX_DX_調査ガイド.html').write_text(htmltext,encoding='utf-8')

doc=Document();sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.65);sec.left_margin=Cm(1.9);sec.right_margin=Cm(1.9)
sec.header_distance=Cm(.65);sec.footer_distance=Cm(.65)
for sty in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    st=doc.styles[sty];st.font.name='Meiryo';st.font.color.rgb=RGBColor(0,0,0);st.font.italic=False
    st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Meiryo')
doc.styles['Normal'].font.size=Pt(10)
doc.styles['Normal'].paragraph_format.line_spacing=Pt(14)
doc.styles['Normal'].paragraph_format.space_after=Pt(6)
for name,size in [('Title',27),('Subtitle',11),('Heading 1',17),('Heading 2',12.5),('Heading 3',11)]:
    doc.styles[name].font.size=Pt(size)
    doc.styles[name].paragraph_format.space_before=Pt(9 if name.startswith('Heading') else 0)
    doc.styles[name].paragraph_format.space_after=Pt(5)
    doc.styles[name].paragraph_format.line_spacing=Pt(size*1.4)
doc.styles['Title'].font.bold=True
hp=sec.header.paragraphs[0];hp.text='JIMTOF2026　AI 自動化 AX DX　事前調査';hp.runs[0].font.size=Pt(8)
fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
fp.add_run('2026年10月19日　｜　').font.size=Pt(8)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');fp._p.append(field)

def p(text,bold=False,size=None):
    pp=doc.add_paragraph();r=pp.add_run(text);r.bold=bold
    if size:r.font.size=Pt(size)
    return pp
def heading(text,level=1):doc.add_heading(text,level)
def page(title,level=1):doc.add_page_break();heading(title,level)
def link(par,label,url):
    rid=par.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rid);r=OxmlElement('w:r');pr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'075E80');pr.append(color);sz=OxmlElement('w:sz');sz.set(qn('w:val'),'16');pr.append(sz);fonts=OxmlElement('w:rFonts');fonts.set(qn('w:eastAsia'),'Meiryo');fonts.set(qn('w:ascii'),'Meiryo');fonts.set(qn('w:hAnsi'),'Meiryo');pr.append(fonts);r.append(pr);t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);par._p.append(h)
def source_line(ids):
    pp=doc.add_paragraph();pp.paragraph_format.space_before=Pt(1);pp.paragraph_format.space_after=Pt(4);pp.paragraph_format.line_spacing=Pt(10)
    for i in dict.fromkeys(ids):
        if i in sources:link(pp,'['+i+'] ',sources[i]['url'])
    for r in pp.runs:r.font.size=Pt(8)
def table(headers,rows,widths=None,size=9):
    tb=doc.add_table(rows=1,cols=len(headers));tb.alignment=WD_TABLE_ALIGNMENT.CENTER;tb.autofit=False
    if widths:
        for i,w in enumerate(widths):tb.columns[i].width=Cm(w)
    for i,h in enumerate(headers):tb.rows[0].cells[i].text=h
    for vals in rows:
        for cell,value in zip(tb.add_row().cells,vals):cell.text=str(value)
    for ri,row in enumerate(tb.rows):
        pr=row._tr.get_or_add_trPr()
        if ri==0:
            repeat=OxmlElement('w:tblHeader');pr.append(repeat)
        nosplit=OxmlElement('w:cantSplit');pr.append(nosplit)
        for ci,cell in enumerate(row.cells):
            if widths:cell.width=Cm(widths[ci])
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props=cell._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'17374A' if ri==0 else ('F1F5F7' if ri%2 else 'FFFFFF'));props.append(shade)
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');margins.append(e)
            props.append(margins);borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
            props.append(borders)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(3);para.paragraph_format.line_spacing=Pt(size*1.35)
                for r in para.runs:r.font.size=Pt(size);r.bold=ri==0;r.font.color.rgb=RGBColor(255,255,255) if ri==0 else RGBColor(0,0,0)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)
    return tb

doc.add_paragraph(d['title'],'Title')
doc.add_paragraph('2024年との比較に基づく技術動向とブース訪問提案','Subtitle')
p('調査基準日 2026年10月19日　公式カタログ取得日 2026年10月16日',size=9)
heading('調査結果の要旨')
p('AIへの言及は、全体と同名企業の比較の両方で増加している。2026年は、加工機の性能に加え、生産準備、段取り、工具、計測、保全をどうつなぐかを比較する価値が高い。フィジカルAIやエージェントの展示は、商用製品と参考出品・提供計画を分けて評価する必要がある。',True)
for i,x in enumerate(d['findings'],1):p(f'{i}　{x}')
table(['全体の調査','比較の軸','視察の提案'],[['両年2,427掲載行を処理','同名538名称キーで比較','大手等33社を必須掲載'],['AI語の全86ヒットを文脈確認','工程と方式の6テーマで分析','優先訪問19件を提示']],[5.7,5.7,5.7],9.5)
heading('読み方と会期',2)
p('本報告書は製造技術・生産技術の視察に向けた事前調査である。2026年の展示内容は予定で、実演の成功や効果を確認したものではない。AXは、AIによって業務の担当・判断・手順が変わる動きを指す整理上の呼称として使う。')
p('会期は10月26〜31日。開場10〜17時、最終日は16時まで。東京ビッグサイト東4〜6ホールは改修で休館。');source_line(['S01'])

page('1 調査方法と情報の充足度')
p('公式出展者一覧の全掲載行を母集団として、両年を独立に集計した。NFKC正規化・空白除去・英字大小統一で名称キーを作り、同名の複数ブースを統合した。完全な法人名寄せではなく、社名変更や別法人の同一性は自動認定していない。')
table(['年度','掲載行','名称キー','説明あり','充足率','説明なし'],coverage_rows,[1.4,2.4,2.7,2.7,2.5,2.7],9)
p('説明ありは、出展説明または「主な出展物」に本文があることを意味する。空欄は関連技術なしとして扱わず、母集団集計は10月16日取得分に固定した。最新発表は個別展示の説明へ反映し、異なる取得時点を混ぜて増減率を作らない。')
heading('全件の把握と深掘りを分ける',2)
p('全件語彙スクリーニング、AI語86ヒットの文脈確認、大手等33社の年度比較、技術方式の横断分析を実施した。既存の334カード・309企業名称の整理も参照したが、重点69カードと追加候補は精査の深さが異なる。カード件数を普及率として使っていない。')
p('主な結論は公式カタログ、企業発表、技術資料を照合した分析である。メーカーの性能・効果の主張は第三者検証の結果ではない。外部資料では2023年の既存機能、2024年の提供開始、2025年の機能追加も確認し、2026年の新発明と取り違えないようにした。')
heading('目次',2)
toc=doc.add_paragraph();f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'TOC \\o "1-1" \\h \\z \\u');toc._p.append(f)

page('2 全体集計と年度比較')
p('両年とも説明文がある538名称キーを比較すると、AI言及は22から40、自動化・省人化は172から193へ増加した。AIは新規言及25、継続15、記載消失7。展示説明における訴求の変化として読む。',True)
doc.add_picture(str(BUILD/'panel.png'),width=Cm(16.8))
doc.paragraphs[-1].paragraph_format.line_spacing=1.0
p('図1 同名538名称キーにおけるテーマ語彙の言及数　出典 保存済み両年CSVを本調査で集計',size=8)
table(['語彙テーマ','2024 件数 比率','2026 件数 比率','同名比較'],[[html.unescape(x).replace('→','から') for x in row] for row in theme_rows],[5.2,4,4,3.8],8.6)
p('比率の分母は説明あり933／800名称キー。語彙区分は重複。AIの語には用途背景やアナログ入出力も含むため、実装件数ではない。説明なし309／331名称キーの状況は不明である。',size=8.5)

page('3 マクロトレンドの読み取り')
heading('AIの訴求は拡大し DXの記載は一様ではない',2)
p('AI言及率は説明あり名称キーで3.1%から7.1%へ増え、同名比較でも増加した。一方、工場運用とデータ連携は同名比較で59から57、設計製造のデジタル化は43から39であり、DX全体が一様に拡大したとは言えない。AIが既存のデジタル基盤の上で、どの仕事を担うようになったかを個別に見る必要がある。')
heading('定量の増加と実装の具体化を照合する',2)
p('AI語の86ヒットを文脈で確認した。2024年のAI/AOモジュールはアナログ入出力、2026年のAIサーバー用材料・AI産業向けスピンドル・翻訳などは、加工や工場でAIを実装する機能の根拠にしない。機械要素やIPCも、応用機能と分ける。')
table(['仕事','2026年の比較対象','読み取れる方向'],[['生産準備','ProPlanAI、CAMAssist、CoAgent、匠フォース','既存の設計・見積・制御環境へAI支援を組み込む'],['加工と段取り','GX Flex、ARMROID、RoboSync、TR-20W','交換・切粉処理等で残る人手を減らす'],['計測と品質','MiSTAR連携、Central、AI摩耗診断','判定結果を補正・交換・継続判断へ返す'],['工場運用','保全支援、遠隔管制、データ基盤','記録から対処の支援へつなぐ']],[3.1,7.1,6.8],9)
p('構造的背景として、2026年版ものづくり白書はAIモデルに加え現場データと実装・運用を重視する政策を整理する。主催者もロボット部品加工やきさげ技能を企画化している。市場全体の導入率や、商用の自律工場が一般化したことの証拠としては使わない。');source_line(['S03','S02','S04'])
heading('比較で残る限界',2)
p('掲載総数の減少は会場構成や掲載定義の差も含む。社名変更・説明の量・展示対象の選択でも数値が変わる。新しい語の追加は新機能の初登場を、記載の消失は機能の廃止を意味しない。')

for block,title in [(d['trends'][:2],'4 生産準備AIとフィジカルAI'),(d['trends'][2:4],'5 自動化と品質管理'),(d['trends'][4:],'6 工場運用と実装基盤')]:
    page(title)
    for t in block:
        heading(t['title'],2);p(t['finding'],True)
        p('これまで　'+t['past']);p('今回　'+t['now'])
        p('比較の軸　'+t['mechanism']+' '+t['implication'])
        p('確認課題　'+t['limit'],size=9)
        p('現地で聞くこと　'+t['question'],True,size=9.5)
        source_line(t['refs'])

groups=defaultdict(list)
for m in d['majors']:groups[m['group']].append(m)
first=True
for group,ms in groups.items():
    batch_size=6 if len(ms)==6 else 5
    for start in range(0,len(ms),batch_size):
        batch=ms[start:start+batch_size]
        page('7 大手と主要企業の比較' if first else group+'の比較',1 if first else 2);first=False
        if start==0:p(group+'。AはAI・自動化の比較を組み立てる優先訪問先、Bは関心工程に応じた追加候補。企業の優劣を示す順位ではない。',size=9)
        for m in batch:
            heading(m['name']+'　'+' '.join(m['booths']),2)
            p('展示予定　'+m['exhibits'],True,size=9.5)
            p('年度比較　'+m['delta'],size=9.5)
            p('確認点　'+m['question'],size=9.5)
            source_line(m['refs'])

for start in range(0,len(d['visits']),7):
    page('8 優先訪問と選定理由' if start==0 else '訪問候補の続き',1 if start==0 else 2)
    if start==0:p('全体を調べたうえで、方式、提供段階、人の仕事の残り方を比較できる19件を選んだ。大手の総合提案と専門企業の具体的なタスクを組み合わせる。全件を1日で回る想定ではない。')
    for v in d['visits'][start:start+7]:
        heading(v['id']+' '+v['name']+'　'+' '.join(v['booths']),2)
        p(v['why']+' 確認することは、'+v['ask'],size=9.5)
        source_line(sorted(v['refs'],key=lambda i: 0 if i.startswith('S') else 1 if '2026' in sources.get(i,{}).get('published','') else 2)[:3])

page('9 訪問計画と現地での確認')
p('以下は時間配分案で、予約や実演時刻ではない。1枠で候補を選び、移動と待ち時間を含める。3社以上を挙げた枠は2社程度、南3も2社程度に絞る。DMG森精機の枠は同社を優先し、ジェイテクトは入替候補。')
for route in d['route']:
    heading(route['day'],2);table(['時間枠','エリア','候補と目的'],route['blocks'],[3,2,12.1],8.5)
p('1日で見る場合　'+d['one_day'],size=9)
p('最終日は16時閉場。東4〜6は休館。実演時刻・混雑・動線に合わせて順番を入れ替える。');source_line(['S01'])

page('10 比較の質問と視察の使い方')
table(['対象','測るべき指標','具体的な質問'],d['questions'],[3.3,6.1,7.7],9.5)
heading('展示段階と実績を分けて記録する',2)
p('商用提供、参考出品、α版、提供計画、動画・サンプル展示を区別する。成功した実演だけでなく、条件を変えた場合、失敗検出、再試行、停止復帰を確認する。価格や市場規模は今回の資料から比較できないため、推定値を置いて順位付けしていない。')
heading('時間指定の候補',2)
p('東ガレリアのKaleido開発者との対談は10月26日14〜15時、10月30日11〜12時と14〜15時が掲載。南4はきさげロボットと匠の協業。ミツトヨの品質管理DXセミナーは10月30日11時20分〜12時05分。参加条件と最新時刻は原典で確認する。');source_line(['S02','S13'])
heading('視察後に判断すること',2)
p('生成・判定・搬送の機能だけでなく、準備と復旧まで含む総工数がどれだけ変わるかを比較する。AIの方式、データの準備負担、検証・承認の位置、他社設備への接続、現場で残る仕事をそろえて記録すると、導入や研究テーマの判断に使いやすい。')
heading('未確認の範囲',2)
p('説明なし企業、2024年の別名照合、会場版の提供条件、実演性能、採用・受注実績、失敗時の復帰は残る確認課題。情報不足を展示なし・進化なしと断定していない。')

for start in [0,14]:
    page('11 主要な出典' if start==0 else '主要な出典の続き',1 if start==0 else 2)
    if start==0:p('個別企業の2024年・2026年公式カタログは各企業のリンクを参照。以下は背景と個別差分を照合した主要資料。その他の技術別根拠はHTMLの出典一覧に収録した。',size=9)
    for s in d['sources'][start:start+14]:
        pp=p('['+s['id']+'] '+s['title'],True,size=9.5)
        link(pp,'　原典を開く',s['url'])
        p(s['locator']+'　公開 '+s['published']+'　確認 '+s['checked'],size=8)
doc.core_properties.title=d['title'];doc.core_properties.subject='2024年と2026年の全出展者一覧を用いたAI 自動化 AX DXの事前調査';doc.core_properties.author='調査レポート'
settings=doc.settings.element;update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
for root_el in [doc.element,doc.styles.element]:
    for b in root_el.xpath('.//w:pBdr'):
        b.getparent().remove(b)
    for pp in root_el.xpath('.//w:pPr'):
        snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');pp.append(snap)
for i in ['TOC 1','TOC 2']:
    if i in doc.styles:
        doc.styles[i].font.name='Meiryo';doc.styles[i].font.size=Pt(9)
        doc.styles[i].paragraph_format.space_after=Pt(2);doc.styles[i].paragraph_format.line_spacing=Pt(12)
doc.save(OUT/'JIMTOF2026_AI_AX_DX_調査報告書.docx')
print(json.dumps({'html_bytes':len(htmltext.encode('utf-8')),'docx':str(OUT/'JIMTOF2026_AI_AX_DX_調査報告書.docx'),'majors':len(d['majors']),'visits':len(d['visits'])},ensure_ascii=False))
